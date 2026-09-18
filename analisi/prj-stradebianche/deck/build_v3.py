"""Aggiunge la vista per legal entity su tre colonne (Cap. 3) e la nota sulla
casistica dimissioni CEO/CdA, a partire dalla versione v2 del deck.

La nuova slide riprende la gabbia della slide gemella «Presidio delle legal
entity nel transitorio» (stesso layout 2_DEFAULT, stessa banda ESEMPIO
ILLUSTRATIVO, titolo bicolore su una riga, sottotitolo a 12 pt).
"""

from pathlib import Path

from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.ns import qn

BLU = RGBColor(0x00, 0x3A, 0x79)
ARANCIO = RGBColor(0xF1, 0x8B, 0x40)
AZZURRO = RGBColor(0xC0, 0xCC, 0xE4)
RIGA_ALT = RGBColor(0xEE, 0xF2, 0xF9)
APERTO = RGBColor(0xFD, 0xEA, 0xD9)   # celle con assegnazione non definita
BIANCO = RGBColor(0xFF, 0xFF, 0xFF)
GRIGIO = RGBColor(0xB3, 0xB3, 0xB3)
FONT = "Century Gothic"

BASE = Path(__file__).resolve().parent
SRC = BASE / "output" / "20260918_SB_Cantiere_Salvezza_vONLINE_v2.pptx"
DST = BASE / "output" / "20260918_SB_Cantiere_Salvezza_vONLINE_v3.pptx"


def style_run(run, size, bold=False, color=BLU):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = FONT
    run.font.color.rgb = color


def box(slide, l, t, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE,
        anchor=MSO_ANCHOR.TOP, margins=(0.15, 0.15, 0.06, 0.06), line_w=0.75):
    sh = slide.shapes.add_shape(shape, Cm(l), Cm(t), Cm(w), Cm(h))
    sh.shadow.inherit = False
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(line_w)
    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left, tf.margin_right = Cm(margins[0]), Cm(margins[1])
    tf.margin_top, tf.margin_bottom = Cm(margins[2]), Cm(margins[3])
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    return sh


def testo(shape, righe, first_align=PP_ALIGN.LEFT):
    """righe = [(testo, dimensione, grassetto, colore), ...]"""
    tf = shape.text_frame
    for i, (txt, size, bold, color) in enumerate(righe):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = first_align
        if i:
            p.space_before = Pt(2)
        style_run(p.add_run(), size, bold, color)
        p.runs[0].text = txt


def riga_mista(shape, spezzoni, align=PP_ALIGN.LEFT):
    """Una sola riga composta da più run con formattazioni diverse."""
    p = shape.text_frame.paragraphs[0]
    p.alignment = align
    for txt, size, bold, color in spezzoni:
        run = p.add_run()
        run.text = txt
        style_run(run, size, bold, color)


def sticker(shape, txt):
    """Marcatore «PER DISCUSSIONE» nello stile della slide gemella:
    testo sottolineato con evidenziatore giallo."""
    p = shape.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    run = p.add_run()
    run.text = txt
    style_run(run, 10, False, BLU)
    rPr = run._r.get_or_add_rPr()
    rPr.set("u", "sng")
    hl = rPr.makeelement(qn("a:highlight"), {})
    clr = hl.makeelement(qn("a:srgbClr"), {"val": "FFFF00"})
    hl.append(clr)
    rPr.append(hl)


prs = Presentation(SRC)
log = []

# ------------------------------------------------------------------ #
# 1. Nota sulla casistica dimissioni CEO/CdA (slide 8)                #
# ------------------------------------------------------------------ #
s08 = prs.slides[7]
nota = box(s08, 1.50, 17.30, 31.00, 1.35, fill=BIANCO, line=ARANCIO,
           shape=MSO_SHAPE.ROUNDED_RECTANGLE, anchor=MSO_ANCHOR.MIDDLE,
           margins=(0.30, 0.30, 0.08, 0.08), line_w=1.0)
testo(nota, [
    ("Casistica in corso di ampliamento", 10, True, BLU),
    ("Oltre ai due estremi qui rappresentati sono possibili scenari intermedi: CdA dimissionario con CEO "
     "in prorogatio; CEO dimissionario con cooptazione non concordata; scenario concordato con cooptazione "
     "del nuovo AD e ricostruzione progressiva del board. Chiusura a valle della call con i legali e della "
     "call dedicata di lunedì", 8.5, False, BLU),
])
log.append("slide 8: nota sull'ampliamento della casistica dimissioni CEO/CdA")

# ------------------------------------------------------------------ #
# 2. Nuova slide: vista per legal entity su tre colonne               #
#    stessa gabbia della slide gemella (slide 18)                     #
# ------------------------------------------------------------------ #
gemella = prs.slides[17]
new = prs.slides.add_slide(gemella.slide_layout)

barra = box(new, 0, 0, 33.87, 0.61, fill=ARANCIO, anchor=MSO_ANCHOR.MIDDLE, margins=(0.2, 0.2, 0, 0))
testo(barra, [("ESEMPIO ILLUSTRATIVO - proposta da validare e discutere con GdL", 11, True, BIANCO)],
      first_align=PP_ALIGN.CENTER)

titolo = box(new, 1.54, 1.12, 30.78, 1.02, margins=(0, 0, 0, 0))
riga_mista(titolo, [
    ("3. Assegnazione perimetri ", 21.4, True, ARANCIO),
    ("| Contenuto dei perimetri e proposta di presidio", 21.4, True, BLU),
])

sott = box(new, 1.54, 2.34, 30.78, 0.53, margins=(0, 0, 0, 0))
testo(sott, [("Presidio da parte di un Chairman e di un responsabile ad hoc per ciascuna LE della target",
              12, False, BLU)])

marc = box(new, 28.40, 2.34, 4.00, 0.53, anchor=MSO_ANCHOR.TOP, margins=(0, 0, 0, 0))
sticker(marc, "PER DISCUSSIONE")

C1_L, C1_W = 1.54, 7.00
C2_L, C2_W = 8.90, 13.10
C3_L, C3_W = 22.35, 10.00

for label, l, w in (("Perimetro / legal entity", C1_L, C1_W),
                    ("Contenuto e attività", C2_L, C2_W),
                    ("Proposta di presidio al Day 1", C3_L, C3_W)):
    hd = box(new, l + 0.16, 3.05, w, 0.55, anchor=MSO_ANCHOR.BOTTOM, margins=(0, 0, 0, 0))
    testo(hd, [(label, 11.5, True, BLU)])

RIGHE = [
    ("Filiali retail", "Rete MPS",
     "Rete commerciale e clientela retail di MPS \u2013 intero perimetro pre carve-out",
     "Banca dei Territori", "S. Barrese", False),
    ("C&IB", "Mediobanca",
     "Attività di corporate & investment banking",
     "IMI CIB", "M. Micillo", False),
    ("WM / PB", "Mediobanca Premier, CMB Monaco, Banca Widiba",
     "Reti di wealth management e private banking. Mediobanca Premier ha ossatura ex CheBanca, "
     "con profilo prevalentemente retail",
     "WMD", "T. Corcos \u2013 conferma in corso", False),
    ("Consumer finance", "Compass e subsidiaries",
     "Credito al consumo",
     "TBD", "Possibile veicolo per il carve-out, in base agli assetti di chiusura OPAS", True),
    ("Leasing / Factoring", "MB Facta, MPS Leasing & Factoring",
     "Leasing e factoring",
     "Banca dei Territori", "S. Barrese", False),
    ("Tech / Ops", "Mediobanca Innovation Services",
     "Servizi tecnologici e operativi infragruppo",
     "CDAITO", "M. Proverbio", False),
    ("AM / Alternatives", "Mediobanca SGR, Polus, RAM",
     "Asset management e gestioni alternative",
     "WMD", "T. Corcos", False),
    ("Fiduciarie / Trust", "MPS Fiduciaria, Spafid",
     "Servizi fiduciari e trust",
     "TBD", "Verificare la sovrapposizione con la fiduciaria già presente nel Gruppo", True),
    ("JV Bancassurance", "AXA-MPS",
     "Joint venture assicurative",
     "WMD", "T. Corcos", False),
]

TOP, PITCH, H = 3.70, 1.34, 1.26

for i, (cluster, entities, descr, assegn, nota_txt, aperto) in enumerate(RIGHE):
    t = TOP + i * PITCH
    box(new, 1.54, t, 30.86, H, fill=(RIGA_ALT if i % 2 else BIANCO))

    c1 = box(new, C1_L + 0.16, t, C1_W, H, anchor=MSO_ANCHOR.MIDDLE, margins=(0, 0.1, 0.02, 0.02))
    testo(c1, [(cluster, 10.5, True, BLU), (entities, 8, False, RGBColor(0x40, 0x6B, 0x9B))])

    c2 = box(new, C2_L, t, C2_W, H, anchor=MSO_ANCHOR.MIDDLE, margins=(0.1, 0.1, 0.02, 0.02))
    testo(c2, [(descr, 9.5, False, BLU)])

    c3 = box(new, C3_L, t + 0.04, C3_W, H - 0.08,
             fill=(APERTO if aperto else BIANCO), line=(ARANCIO if aperto else AZZURRO),
             shape=MSO_SHAPE.ROUNDED_RECTANGLE, anchor=MSO_ANCHOR.MIDDLE,
             margins=(0.18, 0.18, 0.02, 0.02))
    testo(c3, [(assegn, 10, True, BLU), (nota_txt, 8, False, BLU)])

box(new, 1.54, TOP - 0.08, 30.86, 0.02, fill=GRIGIO)

callout = box(new, 1.54, 15.90, 30.86, 1.10, fill=BLU, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
              anchor=MSO_ANCHOR.MIDDLE, margins=(0.35, 0.35, 0.08, 0.08))
testo(callout, [
    ("Principi applicati", 10.5, True, BIANCO),
    ("Assegnazione per prevalenza dell\u2019attività, senza suddividere le società tra più presidi "
     "(no «spezzatino») \u2022 Riduzione delle legal entity come opzione di default \u2022 Per ciascuna "
     "LE resta da identificare il Chairman", 9.5, False, BIANCO),
])

fn = box(new, 1.54, 17.45, 25.60, 1.20, margins=(0, 0, 0, 0))
testo(fn, [
    ("Veicoli di funding e cartolarizzazione e Real Estate non sono inclusi in questa vista: i primi sono "
     "presidiati nell\u2019ambito della Finanza di Gruppo, i secondi sono oggetto di analisi separata data "
     "l\u2019eterogeneità del perimetro (magazzini, tenimenti, sviluppo immobiliare)", 8.5, False, BLU),
    ("A seguire, stessa mappatura per le Aree di Governo: corrispettivo ISP per ciascun responsabile MPS, "
     "da costruire per pivot in assenza di corrispondenza puntuale", 8.5, False, BLU),
])

sld = prs.slides._sldIdLst
el = list(sld)
sld.remove(el[-1])
sld.insert(18, el[-1])
log.append("nuova slide 19: vista per legal entity su tre colonne (dopo la slide gemella sui perimetri)")

prs.save(DST)
print("Salvato:", DST)
print("\n".join(" - " + x for x in log))
