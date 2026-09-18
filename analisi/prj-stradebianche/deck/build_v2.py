"""Genera la slide 'Principi guida' (Cap. 3), la fascia trasversale sulla slide 13
e le correzioni di wording richieste, a partire dal deck originale."""

import copy
from pathlib import Path

from pptx import Presentation
from pptx.util import Cm, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from lxml import etree

NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"

BLU = RGBColor(0x00, 0x3A, 0x79)      # tx2 / accent scuro del deck
ARANCIO = RGBColor(0xF1, 0x8B, 0x40)  # colore titoli
AZZURRO = RGBColor(0xC0, 0xCC, 0xE4)  # accent5
GRIGIO = RGBColor(0xB3, 0xB3, 0xB3)   # linee separatrici
BIANCO = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Century Gothic"

BASE = Path(__file__).resolve().parent
SRC = BASE / "source" / "20260918_SB_Cantiere_Salvezza_vONLINE_a238.pptx"
DST = BASE / "output" / "20260918_SB_Cantiere_Salvezza_vONLINE_v2.pptx"
DST.parent.mkdir(parents=True, exist_ok=True)


def style_run(run, size, bold=False, color=BLU):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = FONT
    run.font.color.rgb = color


def set_bullet(paragraph, color=BLU, mar=200000):
    """Applica il bullet Wingdings usato nel resto del deck."""
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", str(mar))
    pPr.set("indent", str(-mar))
    buClr = etree.SubElement(pPr, NS + "buClr")
    clr = etree.SubElement(buClr, NS + "srgbClr")
    clr.set("val", "%02X%02X%02X" % (color[0], color[1], color[2]))
    buFont = etree.SubElement(pPr, NS + "buFont")
    buFont.set("typeface", "Wingdings")
    buFont.set("pitchFamily", "2")
    buFont.set("charset", "2")
    buChar = etree.SubElement(pPr, NS + "buChar")
    buChar.set("char", "\u00a7")


def add_text(slide, l, t, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE,
             anchor=MSO_ANCHOR.TOP, margins=(0.18, 0.18, 0.10, 0.10), line_w=0.75):
    sh = slide.shapes.add_shape(shape, Cm(l), Cm(t), Cm(w), Cm(h))
    sh.shadow.inherit = False
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(line_w)
    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left, tf.margin_right = Cm(margins[0]), Cm(margins[1])
    tf.margin_top, tf.margin_bottom = Cm(margins[2]), Cm(margins[3])
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    return sh


def add_line(slide, l, t, w, color=GRIGIO):
    from pptx.enum.shapes import MSO_CONNECTOR
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Cm(l), Cm(t), Cm(l + w), Cm(t))
    ln.line.color.rgb = color
    ln.line.width = Pt(0.75)
    return ln


prs = Presentation(SRC)
s02 = prs.slides[1]    # Agenda
s13 = prs.slides[12]   # 2. Pre-Take Control | Attività chiave per il pre-settlement
s16 = prs.slides[15]   # 3. Transitorio | Attività del filone
s17 = prs.slides[16]   # 3. Assegnazione perimetri

log = []

# ---------------------------------------------------------------- #
# 1. Slide 13 - fascia del filone trasversale                       #
# ---------------------------------------------------------------- #
BAND_T, BAND_H = 17.15, 1.55
BAND_R = 27.00          # si ferma prima del logo Intesa Sanpaolo in basso a destra

lbl = add_text(s13, 1.47, BAND_T, 2.92, BAND_H, fill=BLU, anchor=MSO_ANCHOR.MIDDLE,
               margins=(0.10, 0.08, 0.04, 0.04))
p = lbl.text_frame.paragraphs[0]
p.alignment = PP_ALIGN.LEFT
style_run(p.add_run(), 11, True, BIANCO)
p.runs[0].text = "Filone trasversale"

band = add_text(s13, 4.88, BAND_T, BAND_R - 4.88, BAND_H, fill=BIANCO, line=ARANCIO,
                shape=MSO_SHAPE.ROUNDED_RECTANGLE, anchor=MSO_ANCHOR.MIDDLE,
                margins=(0.30, 0.25, 0.08, 0.08), line_w=1.25)
p = band.text_frame.paragraphs[0]
style_run(p.add_run(), 10.5, True, BLU)
p.runs[0].text = ("Action plan sulle prescrizioni regolamentari \u2013 da completare "
                  "entro 6 mesi dal settlement")
p2 = band.text_frame.add_paragraph()
style_run(p2.add_run(), 8.5, False, BLU)
p2.runs[0].text = (
    "Ricognizione di prescrizioni e impegni assunti in sede autorizzativa, con le relative scadenze   |   "
    "Declinazione in attività per ciascun filone, con owner e tempificazione a ritroso   |   "
    "Presidio degli impatti IT, organizzativi e societari   |   "
    "Monitoraggio unitario dell\u2019avanzamento")
log.append("slide 13: aggiunta fascia 'Filone trasversale - action plan prescrizioni regolamentari'")

# ---------------------------------------------------------------- #
# 2. Nuova slide 'Principi guida' (Cap. 3)                          #
# ---------------------------------------------------------------- #
new = prs.slides.add_slide(prs.slides[5].slide_layout)  # stesso layout della slide 6

for ph in list(new.placeholders):
    if ph.placeholder_format.idx == 0:
        ph.left, ph.top, ph.width, ph.height = Cm(1.54), Cm(0.48), Cm(27.30), Cm(1.85)
        ph.text_frame.word_wrap = True
        para = ph.text_frame.paragraphs[0]
        # come nel resto del deck: capitolo in arancio (accent2), titolo in blu
        for testo_run, tema in (("3. Transitorio (pre-/post- CdA) ", MSO_THEME_COLOR.ACCENT_2),
                                ("| Quattro principi guida orientano la gestione del periodo transitorio", None)):
            run = para.add_run()
            run.text = testo_run
            run.font.size = Pt(21.4)
            run.font.name = FONT
            if tema is not None:
                run.font.color.theme_color = tema
    elif ph.placeholder_format.idx == 1:
        ph._element.getparent().remove(ph._element)

sticker = add_text(new, 28.00, 0.60, 4.37, 0.50, anchor=MSO_ANCHOR.TOP,
                   margins=(0, 0, 0, 0))
p = sticker.text_frame.paragraphs[0]
p.alignment = PP_ALIGN.RIGHT
style_run(p.add_run(), 10, False, BLU)
rPr = p.runs[0]._r.get_or_add_rPr()
rPr.set("u", "sng")
hl = etree.SubElement(rPr, NS + "highlight")
etree.SubElement(hl, NS + "srgbClr").set("val", "FFFF00")
p.runs[0].text = "PER DISCUSSIONE"

COL1_L, COL1_W = 2.55, 6.25
COL2_L, COL2_W = 9.30, 10.40
COL3_L, COL3_W = 20.40, 11.90

panel = new.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(20.10), Cm(3.25), Cm(12.40), Cm(12.65))
panel.shadow.inherit = False
panel.fill.solid()
panel.fill.fore_color.rgb = AZZURRO
panel.line.fill.background()
panel._element.getparent().remove(panel._element)
new.shapes._spTree.insert(2, panel._element)  # dietro agli altri elementi

for label, l, w in (("Principio", 1.54, 7.30),
                    ("Cosa vuol dire", COL2_L, COL2_W),
                    ("Implicazioni sulla gestione del transitorio", COL3_L, COL3_W)):
    hd = add_text(new, l, 2.62, w, 0.55, anchor=MSO_ANCHOR.BOTTOM, margins=(0, 0, 0, 0))
    p = hd.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    style_run(p.add_run(), 12, True, BLU)
    p.runs[0].text = label

PRINCIPI = [
    ("PERIMETRO IN USCITA",
     "Minima ingerenza e permanenza sul perimetro in uscita",
     "Su tutto ciò che è destinato al carve-out \u2013 corporate center MPS, filiali e società in uscita \u2013 "
     "l\u2019intervento di ISP è limitato al minimo indispensabile e il tempo di permanenza nel perimetro "
     "del Gruppo è il più breve possibile",
     ["Gestione in continuità, coerente con gli impegni assunti verso Unipol: nessuna manovra elusiva del business",
      "Nessun allineamento ai modelli ISP sul perimetro in uscita (es. politiche di remunerazione, piano di integrazione)",
      "Azioni e comunicazioni al personale segmentate per perimetro di destinazione"]),
    ("PERIMETRO IN INGRESSO",
     "Governance centralizzata sui meccanismi già esistenti",
     "Sugli asset che restano in ISP si applica il modello di indirizzo e coordinamento già in essere nel "
     "Gruppo: non si istituiscono strutture o meccanismi ad hoc per il transitorio",
     ["Attivazione dei riporti funzionali verso i corrispettivi ISP, per Divisione e per Area di Governo",
      "Presidio da rafforzare oltre i soli riporti funzionali, in un contesto atteso meno collaborativo rispetto a UBI",
      "Focus sulle posizioni di governance della target (CEO, CFO, CRO, Comitati), più che sul presidio fisico in loco"]),
    ("ENTRAMBI I PERIMETRI",
     "Riduzione delle legal entity come opzione di default",
     "Coerentemente con l\u2019obiettivo di Piano già comunicato al mercato, l\u2019ipotesi di base è la "
     "razionalizzazione delle legal entity acquisite; il mantenimento è l\u2019eccezione e va motivato",
     ["Verifica delle sovrapposizioni con entità già presenti nel Gruppo (es. fiduciarie/trust)",
      "Le scelte di presidio del transitorio non devono pregiudicare la successiva razionalizzazione societaria"]),
    ("ENTRAMBI I PERIMETRI",
     "Nessuno «spezzatino»: al Day 1 una società, un responsabile",
     "Nel transitorio ciascuna società è assegnata per intero a un unico responsabile, individuato per "
     "prevalenza dell\u2019attività; le entity non vengono suddivise tra più presidi",
     ["L\u2019assegnazione al Day 1 è la migliore ipotesi disponibile ed è rivedibile in una fase successiva",
      "Dove la prevalenza non è univoca l\u2019assegnazione resta TBD (es. Compass, possibile veicolo per il carve-out)"]),
]

ROW_T = [3.40, 6.55, 9.70, 12.85]
ROW_H = 2.90

for i, (tag, nome, definizione, bullets) in enumerate(PRINCIPI):
    top = ROW_T[i]

    badge = add_text(new, 1.54, top + 0.05, 0.80, 0.80, fill=ARANCIO,
                     shape=MSO_SHAPE.OVAL, anchor=MSO_ANCHOR.MIDDLE, margins=(0, 0, 0, 0))
    p = badge.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    style_run(p.add_run(), 12, True, BIANCO)
    p.runs[0].text = str(i + 1)

    box = add_text(new, COL1_L, top, COL1_W, ROW_H, fill=BLU,
                   shape=MSO_SHAPE.ROUNDED_RECTANGLE, anchor=MSO_ANCHOR.MIDDLE,
                   margins=(0.25, 0.20, 0.10, 0.10))
    p = box.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    style_run(p.add_run(), 8, False, AZZURRO)
    p.runs[0].text = tag
    p2 = box.text_frame.add_paragraph()
    p2.alignment = PP_ALIGN.LEFT
    p2.space_before = Pt(4)
    style_run(p2.add_run(), 12, True, BIANCO)
    p2.runs[0].text = nome

    dfn = add_text(new, COL2_L, top, COL2_W, ROW_H, fill=BIANCO, line=AZZURRO,
                   anchor=MSO_ANCHOR.MIDDLE, margins=(0.22, 0.22, 0.12, 0.12))
    p = dfn.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    style_run(p.add_run(), 10, False, BLU)
    p.runs[0].text = definizione

    imp = add_text(new, COL3_L, top, COL3_W, ROW_H, anchor=MSO_ANCHOR.MIDDLE,
                   margins=(0.10, 0.10, 0.08, 0.08))
    tf = imp.text_frame
    for j, b in enumerate(bullets):
        p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
        if j:
            p.space_before = Pt(6)
        p.alignment = PP_ALIGN.LEFT
        set_bullet(p)
        style_run(p.add_run(), 9.5, False, BLU)
        p.runs[0].text = b

    if i:
        add_line(new, 1.54, top - 0.14, 30.90)

banda = add_text(new, 1.54, 16.25, 30.90, 1.75, fill=BLU,
                 shape=MSO_SHAPE.ROUNDED_RECTANGLE, anchor=MSO_ANCHOR.MIDDLE,
                 margins=(0.35, 0.35, 0.12, 0.12))
p = banda.text_frame.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
style_run(p.add_run(), 12, True, BIANCO)
p.runs[0].text = "Chiave di lettura | Definizione nel pre-settlement, attivazione nel transitorio"
p2 = banda.text_frame.add_paragraph()
p2.alignment = PP_ALIGN.LEFT
p2.space_before = Pt(3)
style_run(p2.add_run(), 10.5, False, BIANCO)
p2.runs[0].text = ("Perimetri, riporti funzionali e policy sono definiti prima del settlement; nel periodo "
                   "transitorio se ne attiva l\u2019operatività. Analogamente, il setup della macchina di "
                   "integrazione è impostato nel pre-settlement e attivato nel transitorio")

sld_lst = prs.slides._sldIdLst
elems = list(sld_lst)
sld_lst.remove(elems[-1])
sld_lst.insert(15, elems[-1])
log.append("nuova slide 16: '3. Transitorio | ... quattro principi guida' (inserita dopo il divisorio Cap. 3)")

# ---------------------------------------------------------------- #
# 3. Correzioni di wording                                          #
# ---------------------------------------------------------------- #
def replace_text(shape, new_paragraph_text):
    """Sostituisce il testo mantenendo la formattazione del primo run."""
    para = shape.text_frame.paragraphs[0]
    runs = para.runs
    runs[0].text = new_paragraph_text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


for sh in s16.shapes:
    if not sh.has_text_frame:
        continue
    txt = sh.text_frame.text
    if txt.startswith("Policy:"):
        para = sh.text_frame.paragraphs[0]
        para.runs[0].text = "Policy"
        para.runs[1].text = (": attivazione delle regole standard e di gestione strutturata delle "
                             "eccezioni / casistiche complesse (use cases), definite nel pre-settlement")
        for r in para.runs[2:]:
            r._r.getparent().remove(r._r)
        log.append("slide 17 (ex 16): 'Policy: definizione...' \u2192 'attivazione..., definite nel pre-settlement'")
    elif txt.startswith("Nel corso del periodo transitorio"):
        replace_text(sh, "Il setup della macchina di integrazione è impostato nel pre-settlement "
                         "e attivato nel corso del periodo transitorio")
        log.append("slide 17 (ex 16): nota sul setup della macchina di integrazione spostata al pre-settlement")

for sh in s17.shapes:
    if sh.shape_type != 6:
        continue
    testi = [s.text_frame.text for s in sh.shapes if s.has_text_frame]
    if any(t.startswith("Consumer finance") for t in testi):
        for s in sh.shapes:
            if not s.has_text_frame:
                continue
            if s.text_frame.text == "Banca dei Territori":
                replace_text(s, "TBD")
            elif s.text_frame.text == "S. Barrese":
                replace_text(s, "[…]")
            elif s.text_frame.text.startswith("Incluse"):
                replace_text(s, "Incluse subsidiaries Compass \u2013 possibile veicolo per il carve-out")
        log.append("slide 18 (ex 17): Consumer finance/Compass \u2192 TBD, con nota sul possibile veicolo per il carve-out")
    elif any(t.startswith("Leasing/Factoring") for t in testi):
        for s in sh.shapes:
            if not s.has_text_frame:
                continue
            if s.text_frame.text == "TBD":
                replace_text(s, "Banca dei Territori")
            elif s.text_frame.text == "[…]":
                replace_text(s, "S. Barrese")
        log.append("slide 18 (ex 17): Leasing/Factoring \u2192 Banca dei Territori (S. Barrese)")

for sh in s02.shapes:
    if sh.has_text_frame and sh.text_frame.text.startswith("Assegnazione perimetri"):
        for para in sh.text_frame.paragraphs:          # allineamento definizione/attivazione
            if not para.text.startswith(("Assegnazione", "Policy")):
                continue
            for r in para.runs:
                if r.text == "definizione":
                    r.text = "attivazione"
        src = sh.text_frame.paragraphs[-1]
        nuovo = copy.deepcopy(src._p)
        src._p.addnext(nuovo)
        para = sh.text_frame.paragraphs[-1]
        runs = para.runs
        grassetto = runs[0]
        normale = next(r for r in runs[1:] if not r.font.bold)
        grassetto.text = "Setup della macchina di integrazione"
        normale.text = (": impostazione nel periodo pre-settlement e attivazione nel corso "
                        "del transitorio")
        for r in runs:
            if r is not grassetto and r is not normale:
                r._r.getparent().remove(r._r)
            else:
                r.font.size = Pt(14)
        log.append("slide 2 (Agenda): allineato definizione/attivazione e aggiunto il bullet sul "
                   "setup della macchina di integrazione")
        break

prs.save(DST)
print("Salvato:", DST)
print("\n".join(" - " + x for x in log))
